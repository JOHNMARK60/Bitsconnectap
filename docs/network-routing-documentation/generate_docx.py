"""Generate the formatted Word version of the network documentation."""

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "NETWORK_DOCUMENTATION.md"
OUTPUT = ROOT / "RIPv1-RIPv2_Network_Documentation.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_separate, fld_char_end])


def clean_markdown(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip().replace("  ", " ")


def add_rich_text(paragraph, text: str) -> None:
    """Apply lightweight bold and inline-code formatting."""
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(31, 78, 121)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for level, size, color in [
        (1, 18, RGBColor(31, 78, 121)),
        (2, 14, RGBColor(31, 78, 121)),
        (3, 11, RGBColor(46, 116, 181)),
    ]:
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    if "Code Block" not in [style.name for style in styles]:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(8)
        code_style.paragraph_format.left_indent = Inches(0.2)
        code_style.paragraph_format.right_indent = Inches(0.2)
        code_style.paragraph_format.space_before = Pt(3)
        code_style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "RIPv1–RIPv2 Network Design and Implementation"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    add_field(footer, "PAGE")
    footer.add_run(" of ")
    add_field(footer, "NUMPAGES")
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)


def add_cover(document: Document) -> None:
    for _ in range(3):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("RIPv1–RIPv2 Network Design\nand Implementation Documentation")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Cisco Packet Tracer Dual-Domain Routing Project")
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(89, 89, 89)

    document.add_paragraph()
    details = [
        ("Network 1", "172.16.0.0/21 · Fixed /24 · RIPv1"),
        ("Network 2", "10.10.0.0/24 · VLSM · RIPv2"),
        ("WAN boundary", "10.0.0.0/29 · Two /30 links"),
        ("Date", "July 15, 2026"),
        ("Status", "Implemented and connectivity-verified"),
    ]
    table = document.add_table(rows=len(details), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.7)
    table.columns[1].width = Inches(4.7)
    for row, (label, value) in zip(table.rows, details):
        row.cells[0].text = label
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(row.cells[0], "D9EAF7")
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Prepared as the technical record of addressing, routing, and verification")
    run.italic = True
    run.font.color.rgb = RGBColor(100, 100, 100)
    document.add_page_break()

    heading = document.add_heading("Table of Contents", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    hint = document.add_paragraph("In Microsoft Word, right-click this field and select Update Field to refresh page numbers.")
    hint.runs[0].italic = True
    hint.runs[0].font.size = Pt(8)
    hint.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    document.add_page_break()


def add_table(document: Document, raw_rows: list[str]) -> None:
    rows = []
    for raw in raw_rows:
        cells = [clean_markdown(cell) for cell in raw.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row_index, values in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.text = values[column_index] if column_index < len(values) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(7.5 if column_count >= 7 else 8.5)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if row_index == 0:
                set_cell_shading(cell, "1F4E79")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "EAF2F8")
    set_repeat_table_header(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(document: Document, markdown_line: str) -> None:
    match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", markdown_line.strip())
    if not match:
        return
    alt_text, relative_path = match.groups()
    image_path = ROOT / relative_path
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(7.1))
    if alt_text:
        paragraph.paragraph_format.space_after = Pt(3)


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)


def convert_markdown(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    skipped_source_header = False

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not skipped_source_header:
            if stripped == "---":
                skipped_source_header = True
            index += 1
            continue

        if not stripped or stripped == "---":
            index += 1
            continue

        if stripped.startswith("```"):
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip())
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(document, table_lines)
            continue

        if stripped.startswith("!["):
            add_image(document, stripped)
            index += 1
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            level = min(len(heading_match.group(1)) - 1, 3)
            document.add_heading(clean_markdown(heading_match.group(2)), level=level)
            index += 1
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            add_rich_text(paragraph, bullet_match.group(1))
            index += 1
            continue

        number_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if number_match:
            paragraph = document.add_paragraph(style="List Number")
            add_rich_text(paragraph, number_match.group(1))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate == "---"
                or candidate.startswith("#")
                or candidate.startswith("|")
                or candidate.startswith("![")
                or candidate.startswith("```")
                or re.match(r"^-\s+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        add_rich_text(paragraph, " ".join(paragraph_lines).replace("  ", " "))


def main() -> None:
    document = Document()
    configure_document(document)
    add_cover(document)
    convert_markdown(document, SOURCE.read_text(encoding="utf-8"))

    properties = document.core_properties
    properties.title = "RIPv1–RIPv2 Network Design and Implementation Documentation"
    properties.subject = "Cisco Packet Tracer subnetting, RIP routing, and verification"
    properties.keywords = "RIPv1, RIPv2, VLSM, Cisco Packet Tracer, subnetting"
    properties.comments = "Generated from NETWORK_DOCUMENTATION.md"

    document.save(OUTPUT)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()

